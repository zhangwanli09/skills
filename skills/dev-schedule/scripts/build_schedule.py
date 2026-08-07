#!/usr/bin/env python3
# /// script
# requires-python = ">=3.9"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Build a 开发排程表 (dev scheduling table) xlsx matching the reference format.

Input: a JSON file describing the platform/module/feature/task hierarchy.
Output: an .xlsx with the exact styling of the reference 需求排程表 (blue header,
DengXian 13pt bold header, Microsoft YaHei 11pt body, thin borders, and A/B/C
columns vertically merged to group rows).

Usage:
    python build_schedule.py input.json output.xlsx

Input JSON schema:
{
  "sheet_title": "工作表1",              # optional, default "工作表1"
  "role": "frontend" | "backend",         # optional, informational only
  "platforms": [
    {
      "name": "小程序",                    # column A (服务功能模块)
      "modules": [
        {
          "name": "公共",                  # column B (二级模块)
          "features": [
            {
              "name": "时间筛选组件",       # column C (功能点)
              "tasks": [
                {"name": "时间筛选弹窗", "category": "前端开发", "hours": 1,
                 "api": "", "note": "", "status": "", "owner": ""}
                 # api->E, note->L, status->G, owner->K (all optional)
              ]
            }
          ]
        }
      ]
    }
  ]
}

Each task -> one row. category -> F, hours -> H, api -> E, note -> L.
G (任务状态) defaults to "未开始" unless status is given; I/J left blank;
K comes from owner if present.
Prints a JSON summary (row count, total hours) to stdout.
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADERS = ["服务功能模块", "二级模块", "功能点", "任务名称", "界面调用接口列表",
           "类别", "任务状态", "预计工时（H）", "开始日期", "结束日期",
           "预安排负责人", "备注"]
NCOLS = len(HEADERS)  # 12 -> A..L

# Column widths from the reference (C left at default on purpose).
WIDTHS = {"A": 16.66, "B": 25.66, "D": 60.66, "E": 23.66, "F": 15.66,
          "G": 14.66, "H": 11.66, "I": 18.66, "K": 19.66, "L": 53.66}

HEADER_FILL = PatternFill("solid", fgColor="FF0070C0")
HEADER_FONT = Font(name="等线", size=13, bold=True, color="FFFFFFFF")
BODY_FONT = Font(name="Microsoft YaHei", size=11)
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")

DEFAULT_STATUS = "未开始"
CENTER_COLS = {1, 2, 6, 7, 8}  # A B F G H centered; C/D/rest left


def flatten(platforms):
    """Return (rows, merges). rows: list of 12-value lists. merges: list of
    (col_idx, start_row, end_row) 1-based, only where span > 1."""
    rows = []
    merges = []
    r = 2  # data starts at row 2 (row 1 = header)

    for plat in platforms:
        plat_start = r
        for mod in plat.get("modules", []):
            mod_start = r
            for feat in mod.get("features", []):
                feat_start = r
                tasks = feat.get("tasks", []) or []
                for t in tasks:
                    rows.append([
                        plat.get("name", ""),          # A
                        mod.get("name", ""),           # B
                        feat.get("name", ""),          # C
                        t.get("name", ""),             # D
                        t.get("api", ""),              # E
                        t.get("category", ""),         # F
                        t.get("status", DEFAULT_STATUS),  # G
                        t.get("hours", ""),            # H
                        "", "",                         # I J
                        t.get("owner", ""),            # K
                        t.get("note", ""),             # L
                    ])
                    r += 1
                if r - feat_start > 1:
                    merges.append((3, feat_start, r - 1))  # C
            if r - mod_start > 1:
                merges.append((2, mod_start, r - 1))       # B
        if r - plat_start > 1:
            merges.append((1, plat_start, r - 1))          # A
    return rows, merges


def build(data, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = data.get("sheet_title", "工作表1")

    # Header
    for c, name in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER

    rows, merges = flatten(data.get("platforms", []))

    for i, row in enumerate(rows, start=2):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=i, column=c, value=(val if val != "" else None))
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = CENTER if c in CENTER_COLS else LEFT

    # Merges (A/B/C grouping). Re-assert center-vertical on the anchor cell.
    for col, r1, r2 in merges:
        ws.merge_cells(start_row=r1, start_column=col,
                       end_row=r2, end_column=col)
        anchor = ws.cell(row=r1, column=col)
        anchor.alignment = CENTER if col in CENTER_COLS else LEFT

    # Widths + freeze header
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    wb.save(out_path)

    total = sum(r[7] for r in rows if isinstance(r[7], (int, float)))
    return {"rows": len(rows), "total_hours": total,
            "categories": sorted({r[5] for r in rows if r[5]}),
            "output": out_path}


USAGE = """Usage: build_schedule.py INPUT_JSON OUTPUT_XLSX

Build a 需求排程表 xlsx from a platform/module/feature/task hierarchy.
All styling (blue header, A/B/C vertical merges, column widths, frozen
header) is applied by this script -- do not restyle the output.

Arguments:
  INPUT_JSON   JSON describing the hierarchy; see the schema in this
               file's module docstring.
  OUTPUT_XLSX  Path to write the .xlsx to.

Prints a JSON summary (rows, total_hours, categories, output) to stdout.

Examples:
  python3 scripts/build_schedule.py input.json 需求排程表.xlsx
  uv run scripts/build_schedule.py input.json 需求排程表.xlsx

Exit codes:
  0 success   1 bad arguments   2 invalid input JSON"""


def main():
    args = sys.argv[1:]
    if args and args[0] in ("-h", "--help"):
        print(USAGE)
        sys.exit(0)
    if len(args) < 2:
        print(USAGE, file=sys.stderr)
        sys.exit(1)
    try:
        with open(args[0], encoding="utf-8") as f:
            data = json.load(f)
    except FileNotFoundError:
        print(f"Error: input file not found: {args[0]}", file=sys.stderr)
        sys.exit(2)
    except json.JSONDecodeError as e:
        print(f"Error: {args[0]} is not valid JSON: {e}", file=sys.stderr)
        sys.exit(2)
    summary = build(data, args[1])
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
